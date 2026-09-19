#!/usr/bin/env python3
"""
analyze_deck.py — Stardust deck analyzer.

Called by PHP right after a founder uploads a pitch deck:
    python3 analyze_deck.py /path/to/deck.pdf pdf

Prints a single JSON object to stdout:
    {"excerpt": "...", "keywords": "term1, term2, ...", "error": null}

Designed to degrade gracefully: if pdfplumber / python-pptx aren't
installed, or the file can't be parsed, it still returns valid JSON
(with an "error" note) instead of crashing the PHP caller.
"""

import sys
import json
import re
from collections import Counter

STOPWORDS = set("""
a an the and or but if then else for of to in on at by with as is are was
were be been being this that these those it its it's we our you your they
their he she his her from into over under about above below up down out
not no nor so than too very can will just also more most such only own
company startup startups product products team market users user customer
customers business plan slide deck pitch page
""".split())

MIN_WORD_LEN = 4
MAX_EXCERPT_CHARS = 900
TOP_KEYWORDS = 12


def extract_pdf_text(path):
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(path) as pdf:
            for pg in pdf.pages[:12]:
                t = pg.extract_text()
                if t:
                    text_parts.append(t)
        return "\n".join(text_parts), None
    except ImportError:
        pass
    except Exception as e:
        return "", f"pdfplumber failed: {e}"

    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(path)
        text_parts = [p.extract_text() or "" for p in reader.pages[:12]]
        return "\n".join(text_parts), None
    except ImportError:
        return "", "No PDF text library installed (pip install pdfplumber)"
    except Exception as e:
        return "", f"PyPDF2 failed: {e}"


def extract_pptx_text(path):
    try:
        from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)
        return "\n".join(parts), None
    except ImportError:
        return "", "python-pptx not installed (pip install python-pptx)"
    except Exception as e:
        return "", f"python-pptx failed: {e}"


def top_keywords(text):
    words = re.findall(r"[A-Za-z][A-Za-z\-]{2,}", text.lower())
    words = [w for w in words if len(w) >= MIN_WORD_LEN and w not in STOPWORDS]
    counts = Counter(words)
    return [w for w, _ in counts.most_common(TOP_KEYWORDS)]


def main():
    if len(sys.argv) < 3:
        print(json.dumps({"excerpt": "", "keywords": "", "error": "usage: analyze_deck.py <path> <pdf|pptx|ppt>"}))
        return

    path, ext = sys.argv[1], sys.argv[2].lower()

    if ext == "pdf":
        text, err = extract_pdf_text(path)
    elif ext in ("pptx", "ppt"):
        text, err = extract_pptx_text(path)
    else:
        text, err = "", f"Unsupported file type: {ext}"

    text = re.sub(r"\s+", " ", text or "").strip()
    excerpt = text[:MAX_EXCERPT_CHARS]
    keywords = top_keywords(text) if text else []

    print(json.dumps({
        "excerpt": excerpt,
        "keywords": ", ".join(keywords),
        "error": err,
    }))


if __name__ == "__main__":
    main()
